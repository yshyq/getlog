from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "Kubernetes节点日志下载门户开发文档-v1.1.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
GRAY = "555555"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "000000"
RED = "9B1C1C"
GOLD = "7A5A00"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(
    run,
    *,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = latin
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in: Sequence[float], indent_dxa: int = 120):
    widths_dxa = [round(width * 1440) for width in widths_in]
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            cell.width = Inches(widths_in[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_keep(paragraph, *, next_: bool = False, lines: bool = False):
    p_pr = paragraph._p.get_or_add_pPr()
    if next_:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)
    if lines:
        keep_lines = OxmlElement("w:keepLines")
        p_pr.append(keep_lines)


def set_paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=GRAY)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = rgb(NAVY)
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(GRAY)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Kubernetes 节点日志下载门户  /  开发文档")
    set_run_font(run, size=8.5, color=GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    add_page_field(p)
    run = p.add_run(" 页")
    set_run_font(run, size=9, color=GRAY)


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    set_paragraph_keep(p, next_=True, lines=True)
    return p


def add_para(
    doc: Document,
    text: str = "",
    *,
    bold_prefix: str | None = None,
    italic: bool = False,
    color: str | None = None,
    align=None,
    after: float | None = None,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color)
    return p


def add_bullets(doc: Document, items: Iterable[str], *, level: int = 0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbers(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_code(doc: Document, code: str):
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, CALLOUT)
    p.paragraph_format.keep_together = True
    for index, line in enumerate(code.strip("\n").splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, latin="Consolas", size=8.5, color=NAVY)
    return p


def add_callout(doc: Document, title: str, text: str, *, fill: str = CALLOUT, accent: str = BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    set_paragraph_shading(p, fill)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r = p.add_run(f"{title}：")
    set_run_font(r, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    set_paragraph_keep(p, lines=True)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float],
    *,
    compact: bool = False,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9 if compact else 9.5, bold=True, color=NAVY)

    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=8.5 if compact else 9.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover(doc: Document):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    r = kicker.add_run("TECHNICAL DEVELOPMENT GUIDE")
    set_run_font(r, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("Kubernetes 节点日志下载门户")
    set_run_font(r, size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    r = subtitle.add_run("开发文档")
    set_run_font(r, size=18, color=DARK_BLUE)

    add_table(
        doc,
        ["文档项", "内容"],
        [
            ["版本", "v1.1-r2"],
            ["日期", "2026-07-10"],
            ["状态", "开发基线（评审修复版）"],
            ["适用范围", "第一版：单共享管理员、单集群、只读日志浏览与下载"],
            ["设计依据", "docs/superpowers/specs/2026-07-09-kubernetes-log-download-portal-design.md"],
        ],
        [1.5, 5.0],
    )

    add_callout(
        doc,
        "开发结论",
        "采用 Go 中央门户 + Nginx DaemonSet。门户负责认证、节点发现、参数校验、审计与流式代理；Nginx 仅在集群内部只读列举和传输所在节点的日志文件。",
        fill=CALLOUT,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run("面向后端、前端、测试、平台工程与安全评审人员")
    set_run_font(r, size=9.5, color=GRAY, italic=True)
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)

    add_heading(doc, "1. 文档说明", 1)
    add_para(
        doc,
        "本文把已确认的产品设计展开为可直接实施的开发基线，规定工程边界、模块职责、接口契约、安全控制、部署资源、测试方法与交付顺序。若实现与本文冲突，以安全边界和验收标准为优先，并通过设计变更记录统一修订。",
    )

    add_heading(doc, "1.1 目标与读者", 2)
    add_bullets(
        doc,
        [
            "后端开发：据此创建 Go 工程、实现 API、Kubernetes 发现、Nginx 客户端和下载代理。",
            "前端开发：据此实现登录、节点与服务选择、文件列表、过滤、刷新和错误状态。",
            "平台工程：据此准备 Deployment、DaemonSet、Ingress、RBAC、NetworkPolicy、ConfigMap 与 Secret。",
            "测试与安全：据此建立功能、攻击面、性能、部署和验收用例。",
        ],
    )

    add_heading(doc, "1.2 第一版边界", 2)
    add_para(doc, "第一版仅解决“节点 → 服务 → 文件”的安全只读下载，不承担日志平台职责。")
    add_table(
        doc,
        ["包含", "明确不包含"],
        [
            ["共享管理员登录与退出", "多用户、角色、服务级授权"],
            ["Node 状态、服务列表、文件元数据", "全文搜索、实时 Tail、内容预览"],
            ["单文件流式下载与单区间 Range", "多文件打包、上传、编辑、删除"],
            ["标准输出 JSON 审计", "业务数据库、日志聚合与告警"],
            ["单 Kubernetes 集群、单门户副本", "跨集群、门户高可用"],
        ],
        [3.25, 3.25],
    )

    add_heading(doc, "1.3 实现假设", 2)
    add_bullets(
        doc,
        [
            "所有目标 Node 的宿主机日志根目录相同，服务目录是根目录的直接子目录，服务目录内不递归。",
            "门户以一个 Go 二进制交付，前端静态资源通过 embed 内嵌；前后端同源，不引入独立前端运行时。",
            "门户直接访问 Ready Nginx Pod IP；普通 Service 仅用于门户自身，不用于节点日志转发。",
            "会话采用服务端内存存储，浏览器 Cookie 仅保存不可预测的会话标识及签名。门户重启后会话失效。",
            "集群已有 Ingress Controller、CNI NetworkPolicy 能力和统一日志采集能力。",
        ],
    )

    add_heading(doc, "2. 总体架构与调用链", 1)
    add_code(
        doc,
        """
用户浏览器
    │ HTTPS
    ▼
Ingress ──► Portal Service ──► Go Portal Pod
                               ├─ watch/list ─► Kubernetes API
                               ├─ HTTP GET ───► Node A Nginx Pod ─► /logs (hostPath, ro)
                               └─ HTTP GET ───► Node B Nginx Pod ─► /logs (hostPath, ro)
""",
    )
    add_para(
        doc,
        "浏览器始终只连接门户。门户从 informer 缓存中把逻辑节点名解析为当前 Ready、非终止状态的 Nginx Pod IP；文件列表与下载请求再发往该 Pod。下载正文通过固定大小缓冲区边读边写，门户不落盘、不完整缓冲。",
    )

    add_heading(doc, "2.1 组件职责", 2)
    add_table(
        doc,
        ["组件", "职责", "禁止事项"],
        [
            ["浏览器 UI", "登录、选择、展示、过滤、触发下载", "不接收真实路径、Pod IP 或内部 URL"],
            ["Go Portal", "认证、会话、配置、节点发现、参数校验、审计、流式代理", "不写日志文件、不缓存完整文件"],
            ["Kubernetes API", "提供 Node 与 DaemonSet Pod 状态", "不向浏览器暴露"],
            ["Nginx DaemonSet", "JSON autoindex、静态文件、Range、ETag、Last-Modified", "无外部 Ingress/NodePort，不访问 K8s API"],
            ["宿主机日志目录", "保存业务日志", "仅以只读 hostPath 挂载"],
        ],
        [1.25, 3.1, 2.15],
        compact=True,
    )

    add_heading(doc, "2.2 核心时序", 2)
    add_numbers(
        doc,
        [
            "登录：验证凭据 → 建立内存会话 → 写入安全 Cookie → 记录 login 审计事件。",
            "浏览：读取节点快照 → 选择节点与服务 → 解析 Ready Pod → 请求 Nginx autoindex JSON → 过滤普通文件 → 返回 DTO。",
            "下载：复核会话和全部逻辑参数 → 解析 Ready Pod → 校验 Range → 建立上游请求 → 复制允许的响应头和响应流 → 记录 download 审计事件。",
            "异常：生成 requestId → 对内记录根因 → 对外返回稳定错误码和非敏感消息。",
        ],
    )

    add_heading(doc, "3. 技术选型与工程结构", 1)
    add_table(
        doc,
        ["层面", "推荐实现"],
        [
            ["语言与运行时", "Go；版本在 go.mod 与构建镜像中固定，使用当前团队批准的稳定版本"],
            ["HTTP", "标准 net/http；路由器可选轻量依赖，但必须保留原始 Request Context"],
            ["Kubernetes", "client-go informer/lister，限制到 Node 和指定 namespace/label 的 Pod"],
            ["密码", "bcrypt 或 Argon2id；实现只选择一种并通过配置迁移，不并行自研算法"],
            ["前端", "原生 TypeScript/轻量构建或 Go 模板 + 少量 JS；编译产物使用 go:embed"],
            ["日志", "结构化 JSON 写 stdout；运行日志与审计事件使用不同 eventType"],
            ["测试", "Go testing、httptest、fake clientset、测试 Nginx 容器、Kubernetes 测试环境"],
        ],
        [1.45, 5.05],
    )

    add_heading(doc, "3.1 推荐仓库结构", 2)
    add_code(
        doc,
        """
.
├─ cmd/portal/main.go
├─ internal/
│  ├─ app/app.go
│  ├─ config/{config.go,validate.go}
│  ├─ auth/{password.go,session.go,ratelimit.go,middleware.go}
│  ├─ api/{router.go,response.go,auth_handler.go,node_handler.go,file_handler.go}
│  ├─ discovery/{informer.go,snapshot.go,resolver.go}
│  ├─ agent/{client.go,autoindex.go}
│  ├─ download/{proxy.go,range.go,headers.go}
│  ├─ security/{component.go,origin.go,client_ip.go}
│  ├─ audit/{event.go,logger.go}
│  └─ webui/{handler.go,dist/}
├─ web/{src/,package.json}
├─ deploy/base/
│  ├─ portal-{configmap,secret,deployment,service,ingress}.yaml
│  ├─ portal-{serviceaccount,role,rolebinding}.yaml
│  ├─ agent-{configmap,daemonset,networkpolicy}.yaml
│  └─ kustomization.yaml
├─ test/{integration,e2e}/
├─ Dockerfile
├─ Makefile
├─ go.mod
└─ README.md
""",
    )
    add_callout(
        doc,
        "边界规则",
        "api 包只处理 HTTP 适配；auth、discovery、agent、download、security 和 audit 均通过小接口注入。禁止 handler 直接调用 Kubernetes Client 或拼接上游 URL。",
    )

    add_heading(doc, "4. 配置与密钥", 1)
    add_heading(doc, "4.1 ConfigMap 配置模型", 2)
    add_code(
        doc,
        """
logRoot: /data/logs
services:
  - id: order-service
    displayName: 订单服务
    directory: order-service
sessionTTL: 8h
discovery:
  namespace: log-portal
  podSelector: app.kubernetes.io/name=node-log-agent
  staleAfter: 30s
  watchRefreshInterval: 20s
upstream:
  port: 8080
  dialTimeout: 3s
  responseHeaderTimeout: 10s
fileList:
  maxItems: 10000
  maxResponseBytes: 8MiB
download:
  maxConcurrent: 8
  bufferSize: 64KiB
  bodyIdleTimeout: 60s
  maxDuration: 0
agent:
  probeFileName: .log-portal-read-probe
  probeScope: allServices
auth:
  maxFailures: 5
  failureWindow: 5m
  blockDuration: 15m
""",
    )
    add_table(
        doc,
        ["字段", "校验规则", "默认/行为"],
        [
            ["logRoot", "绝对路径；clean 后不为 /；部署值与 hostPath 相同", "必填"],
            ["services[].id", "唯一；^[a-z0-9][a-z0-9-]{0,62}$", "浏览器逻辑标识"],
            ["services[].displayName", "去首尾空白后 1–80 个字符", "必填"],
            ["services[].directory", "精确匹配 ^[a-z0-9][a-z0-9_-]{0,62}$；不得含 .、..、/、\\、%", "必填"],
            ["sessionTTL", "15m–24h", "8h"],
            ["discovery.staleAfter", "5s–5m", "30s"],
            ["discovery.watchRefreshInterval", "5s–staleAfter，强制周期性重建 watch 或 relist", "20s"],
            ["fileList.maxItems", "1–10000；流式解码到第 N+1 项立即终止", "10000"],
            ["fileList.maxResponseBytes", "1MiB–64MiB；读取前施加硬上限", "8MiB"],
            ["download.maxConcurrent", "1–128", "8"],
            ["download.bufferSize", "16KiB–1MiB", "64KiB"],
            ["download.bodyIdleTimeout", "5s–10m；上游或客户端连续无字节进展即取消", "60s"],
            ["download.maxDuration", "0 或 1m–24h；0 表示不设置总下载时长", "0"],
            ["agent.probeFileName", "单一安全文件名；必须存在于每个允许服务目录", ".log-portal-read-probe"],
        ],
        [1.65, 3.4, 1.45],
        compact=True,
    )

    add_heading(doc, "4.2 Secret 与环境变量", 2)
    add_table(
        doc,
        ["环境变量", "来源", "要求"],
        [
            ["PORTAL_USERNAME", "Secret", "固定管理员用户名，建议 admin；不得写入 ConfigMap"],
            ["PORTAL_PASSWORD_HASH", "Secret", "bcrypt 或 Argon2id 编码哈希；拒绝明文"],
            ["PORTAL_SESSION_KEY", "Secret", "至少 32 字节随机值，Base64 编码；用于 Cookie HMAC"],
            ["PORTAL_CONFIG_PATH", "Deployment", "配置文件挂载路径"],
            ["PORTAL_TRUSTED_PROXY_CIDRS", "ConfigMap/Deployment", "仅信任 Ingress 网段，其他来源忽略 X-Forwarded-For"],
        ],
        [2.15, 1.55, 2.8],
        compact=True,
    )
    add_callout(
        doc,
        "启动失败策略",
        "配置解析、目录白名单、密钥长度、重复服务 ID 或超出安全范围时，进程必须在监听端口前退出，并输出不含密钥值的结构化错误。",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "5. 核心领域模型与内部接口", 1)
    add_code(
        doc,
        """
type NodeStatus string
const (
    NodeAvailable        NodeStatus = "READY"
    NodeNotReady         NodeStatus = "NODE_NOT_READY"
    NodeAgentUnavailable NodeStatus = "AGENT_UNAVAILABLE"
)

type NodeView struct {
    Name       string     `json:"name"`
    NodeReady  bool       `json:"nodeReady"`
    AgentReady bool       `json:"agentReady"`
    Status     NodeStatus `json:"status"`
    Reason     string     `json:"reason,omitempty"`
}

type ServiceView struct {
    ID          string `json:"id"`
    DisplayName string `json:"displayName"`
}

type FileView struct {
    Name         string    `json:"name"`
    Size         int64     `json:"size"`
    LastModified time.Time `json:"lastModified"`
}
""",
    )
    add_heading(doc, "5.1 依赖接口", 2)
    add_code(
        doc,
        """
type NodeResolver interface {
    ListNodes(ctx context.Context) ([]NodeView, error)
    ResolveReadyAgent(ctx context.Context, node string) (AgentEndpoint, error)
}

type AgentClient interface {
    ListFiles(ctx context.Context, endpoint AgentEndpoint, directory string) ([]FileMeta, error)
    OpenFile(ctx context.Context, endpoint AgentEndpoint, directory, filename string,
        headers http.Header) (*http.Response, error)
}

type SessionStore interface {
    Create(username string, ttl time.Duration) (Session, error)
    Get(id string) (Session, bool)
    Delete(id string)
}

type AuditLogger interface {
    Login(LoginEvent)
    Download(DownloadEvent)
    Security(SecurityEvent)
}
""",
    )
    add_para(
        doc,
        "接口返回领域错误而非 HTTP 状态；api 层集中完成状态映射。测试可以使用内存实现或 httptest Server，避免 handler 测试依赖真实集群。",
    )

    add_heading(doc, "6. API 通用约定", 1)
    add_bullets(
        doc,
        [
            "所有 API 使用 /api/v1 前缀；除 login 外均要求有效会话。",
            "JSON 响应使用 UTF-8；成功列表使用对象包裹，便于后续增加分页或元数据。",
            "每个请求接收或生成 X-Request-ID；仅接受 ^[A-Za-z0-9_-]{1,64}$ 的客户端值，缺失或非法时生成新的 UUID/ULID；响应始终回传最终值，日志与审计均记录同一值。",
            "禁止在响应中出现 Pod IP、宿主机绝对路径、上游 URL、堆栈、密钥或原始底层错误。",
            "所有非下载错误使用统一 envelope；下载在响应头写出后发生的错误只能中断连接并记录审计。",
            "动态数据响应设置 Cache-Control: no-store；静态前端资源使用内容哈希与长期缓存。",
        ],
    )
    add_code(
        doc,
        """
{
  "error": {
    "code": "AGENT_UNAVAILABLE",
    "message": "节点日志代理不可用",
    "requestId": "01J..."
  }
}
""",
    )
    add_table(
        doc,
        ["HTTP", "错误码", "适用场景"],
        [
            ["400", "INVALID_ARGUMENT", "参数非法、Range 非法、路径保护命中"],
            ["401", "UNAUTHENTICATED", "未登录、Cookie 无效或会话过期"],
            ["404", "NODE_NOT_FOUND / SERVICE_NOT_FOUND / FILE_NOT_FOUND", "对象不存在；服务白名单外也返回 404"],
            ["429", "RATE_LIMITED / DOWNLOAD_LIMIT_REACHED", "登录频率或下载并发超过限制"],
            ["413", "FILE_LIST_TOO_LARGE", "autoindex 响应体或文件项数量超过配置上限"],
            ["502", "UPSTREAM_ERROR", "Agent 返回 403、非预期 5xx、重定向、畸形 JSON 或非法响应头"],
            ["503", "AGENT_UNAVAILABLE / KUBERNETES_UNAVAILABLE", "节点无 Ready Agent 或发现缓存不可用"],
            ["504", "UPSTREAM_TIMEOUT", "连接 Nginx 或等待响应头超时"],
            ["500", "INTERNAL", "未分类服务器错误"],
        ],
        [0.7, 2.25, 3.55],
        compact=True,
    )

    add_heading(doc, "7. API 详细契约", 1)
    add_heading(doc, "7.1 POST /api/v1/login", 2)
    add_para(doc, "请求体：")
    add_code(doc, '{"username":"admin","password":"<用户输入>"}')
    add_para(
        doc,
        "成功返回 204，并设置会话 Cookie。失败凭据统一返回 401，不区分用户名不存在或密码错误。被限流返回 429 和 Retry-After。请求体上限 8 KiB，拒绝额外字段可在严格模式下启用。",
    )

    add_heading(doc, "7.2 POST /api/v1/logout", 2)
    add_para(
        doc,
        "要求有效会话和同源 Origin/Host 校验。删除服务端会话并下发过期 Cookie，返回 204。重复退出保持幂等；无有效会话时仍可清理 Cookie，但对 API 调用返回 401。",
    )

    add_heading(doc, "7.3 GET /api/v1/session", 2)
    add_code(
        doc,
        """
{
  "authenticated": true,
  "username": "admin",
  "expiresAt": "2026-07-09T18:00:00+08:00"
}
""",
    )

    add_heading(doc, "7.4 GET /api/v1/nodes", 2)
    add_code(
        doc,
        """
{
  "items": [
    {"name":"worker-a","nodeReady":true,"agentReady":true,"status":"READY"},
    {"name":"worker-b","nodeReady":false,"agentReady":true,"status":"NODE_NOT_READY","reason":"node_not_ready"},
    {"name":"worker-c","nodeReady":true,"agentReady":false,"status":"AGENT_UNAVAILABLE","reason":"no_ready_agent_pod"}
  ],
  "observedAt": "2026-07-09T10:00:00+08:00"
}
""",
    )
    add_para(
        doc,
        "节点按名称升序返回。nodeReady 由 NodeCondition 判断；agentReady 由同节点非终止且 Ready 的 DaemonSet Pod 判断。status 是派生展示字段：agentReady=false 时一律为 AGENT_UNAVAILABLE 并禁止文件列表/下载；agentReady=true 且 nodeReady=false 时为 NODE_NOT_READY，仅作为风险提示，仍允许尝试读取日志；两者均为 true 时为 READY。缓存 observedAt 仅在成功完成 list/relist 并提交快照，或 watch 处理到事件/BOOKMARK 并推进 resourceVersion 后更新；仅建立 watch 连接但未推进 resourceVersion 不刷新新鲜度。超过 staleAfter 时返回 503，不以陈旧 Pod IP 发起新请求。",
    )

    add_heading(doc, "7.5 GET /api/v1/nodes/{node}/services", 2)
    add_code(
        doc,
        """
{
  "node":"worker-a",
  "agentReady":true,
  "items":[
    {"id":"order-service","displayName":"订单服务"},
    {"id":"payment-service","displayName":"支付服务"}
  ]
}
""",
    )
    add_para(
        doc,
        "服务列表来自已校验 ConfigMap，不读取宿主机目录。节点必须存在；Agent 暂时不可用时仍可返回服务列表，但同时返回 agentReady=false，前端禁用文件加载。",
    )

    add_heading(doc, "7.6 GET /api/v1/nodes/{node}/services/{service}/files", 2)
    add_code(
        doc,
        """
{
  "node":"worker-a",
  "service":"order-service",
  "items":[
    {"name":"app.log","size":1048576,"lastModified":"2026-07-09T01:20:30Z"}
  ],
  "observedAt":"2026-07-09T02:00:00Z"
}
""",
    )
    add_para(
        doc,
        "门户在读取前对 Nginx autoindex 响应体施加 maxResponseBytes 硬上限，并使用 json.Decoder 流式解析数组；仅保留 type=file、名称通过安全校验、size≥0 且 mtime 可解析的项。达到第 maxItems+1 项、响应体超限或 JSON 非法时立即取消上游请求，返回 413 FILE_LIST_TOO_LARGE 或 502 UPSTREAM_ERROR，不得先完整载入内存。结果按最后修改时间倒序、文件名升序稳定排序。",
    )

    add_heading(doc, "7.7 GET /api/v1/nodes/{node}/services/{service}/files/{filename}/download", 2)
    add_bullets(
        doc,
        [
            "请求允许 Range 与 If-Range；第一版只接受单一区间，拒绝逗号分隔的多区间。",
            "上游请求固定设置 Accept-Encoding: identity，独立 http.Transport 设置 DisableCompression=true，保持 Range、Content-Length 与 Content-Range 一致。",
            "未带 Range 的请求只接受上游 200；带 Range 的请求只接受 206 或 416。206/416 必须带合法 Content-Range；Content-Length 必须为非负十进制且与状态语义一致，否则按非法响应头处理。",
            "上游 403、3xx、非预期 5xx、Range 请求返回 200 或非法响应头统一映射为 502 UPSTREAM_ERROR，不向浏览器暴露 Pod IP 或底层响应正文。",
            "响应固定设置 Content-Disposition: attachment; filename*=UTF-8''<escaped>。",
            "仅转发 Content-Type、Content-Length、Content-Range、Accept-Ranges、ETag、Last-Modified。",
            "不复制 Connection、Transfer-Encoding、Set-Cookie、Server 或上游 Content-Disposition。",
            "客户端取消后立即取消上游 context，释放并发槽并记录 completed=false。",
        ],
    )

    add_heading(doc, "8. 认证、会话与登录限流", 1)
    add_heading(doc, "8.1 会话设计", 2)
    add_numbers(
        doc,
        [
            "使用 crypto/rand 生成 32 字节 session ID，服务端内存表保存用户名、签发时间、过期时间。",
            "Cookie 值为 base64url(sessionID).base64url(HMAC-SHA256)，校验使用 hmac.Equal。",
            "每次受保护请求校验签名、会话存在性和 expiresAt；过期会话惰性删除，后台定期清理。",
            "Cookie 属性：HttpOnly、Secure、Path=/、SameSite=Lax、Max-Age=sessionTTL；不设置 Domain。",
            "登录成功旋转会话；退出立即删除。门户重启清空内存会话，符合第一版约束。",
        ],
    )
    add_heading(doc, "8.2 密码与限流", 2)
    add_bullets(
        doc,
        [
            "密码哈希参数在启动时验证；不接受未知前缀、明文或过低成本参数。",
            "用户名不存在时执行等成本的虚拟哈希校验，降低计时差异。",
            "按可信客户端 IP 维护滑动窗口失败计数，并设置全局登录速率上限，防止单 IP 绕过。",
            "成功登录清除该 IP 失败计数；失败原因只在内部审计分类，对外统一为“用户名或密码错误”。",
            "POST 接口检查 Origin；无 Origin 的非浏览器客户端按部署策略允许或拒绝，默认要求 Origin 与外部 Host 一致。",
        ],
    )

    add_heading(doc, "9. Kubernetes 节点与 Agent 发现", 1)
    add_heading(doc, "9.1 Informer 范围", 2)
    add_bullets(
        doc,
        [
            "Node informer：获取/list/watch Node，仅读取名称与 Ready Condition。",
            "Pod informer：限制到 agent namespace 和固定 label selector，仅读取 spec.nodeName、status.podIP、Ready Condition、deletionTimestamp 与 creationTimestamp。",
            "启动必须等待两个 informer cache sync；超时则 readiness=false，但进程可继续重试。",
            "snapshot 以不可变结构原子替换，HTTP 热路径只读，不持有 informer 锁。",
        ],
    )

    add_heading(doc, "9.2 节点到 Pod 的解析规则", 2)
    add_numbers(
        doc,
        [
            "使用 Kubernetes validation.IsDNS1123Subdomain 校验 node（最大 253 字节，可包含点分段），并在 Node 快照中精确匹配；不得误用最大 63 字节的 DNS label 规则。",
            "筛选 spec.nodeName 等于目标、Pod Ready=true、podIP 非空且 deletionTimestamp 为空的 Agent Pod。",
            "无候选返回 AGENT_UNAVAILABLE；多候选时选择 creationTimestamp 最新者，并输出 duplicate_agent_pods 告警。",
            "构建 http://<podIP>:<configuredPort>；IPv6 地址使用 net.JoinHostPort，不手工拼接冒号。",
            "每次文件列表或下载都重新解析；不把上次 Pod IP 存入浏览器或长时间业务缓存。",
        ],
    )
    add_callout(
        doc,
        "Kubernetes 故障",
        "watch 断开由 informer 自动重连。新鲜度时间仅在成功完成 list/relist 并提交快照，或 watch 处理到事件/BOOKMARK 且推进 resourceVersion 后更新；单纯建立 watch 连接、未收到 BOOKMARK/事件或未推进 resourceVersion 时不得刷新新鲜度。由于 BOOKMARK 不保证固定间隔，watch 使用有界超时并按 watchRefreshInterval 周期性重建；若窗口内没有事件或 BOOKMARK，必须主动 relist 来证明快照仍可用。超过 staleAfter 时拒绝新列表和下载请求并返回 503。已经建立的下载流不依赖 Kubernetes API，不主动中断。",
    )

    add_heading(doc, "10. Nginx DaemonSet 实现", 1)
    add_heading(doc, "10.1 Nginx 配置基线", 2)
    add_code(
        doc,
        """
server {
    listen 8080;
    server_tokens off;

    location /logs/ {
        alias /logs/;
        autoindex on;
        autoindex_format json;
        autoindex_localtime off;
        disable_symlinks on from=/logs;
        sendfile on;
        limit_except GET { deny all; }
    }
}
""",
    )
    add_para(
        doc,
        "门户仅访问 /logs/{escapedDirectory}/ 和 /logs/{escapedDirectory}/{escapedFilename}。directory 必须来自服务白名单，filename 必须作为单个路径段独立转义。不得把用户提供的完整路径直接附加到上游 URL。",
    )

    add_heading(doc, "10.2 Pod 安全上下文", 2)
    add_bullets(
        doc,
        [
            "runAsNonRoot=true，固定 runAsUser/runAsGroup，allowPrivilegeEscalation=false。",
            "readOnlyRootFilesystem=true，drop: [ALL]；仅为 /var/cache/nginx、/var/run 等必要目录挂载 emptyDir。",
            "hostPath 挂载到 /logs，readOnly=true；hostPath type 使用 Directory。",
            "部署前验证宿主机目录对 Agent 的固定 UID/GID 或 supplementalGroups 具备目录遍历和文件读取权限；hostPath 只读不等于可读，且不得依赖 fsGroup 修改宿主机所有权。",
            "Agent readiness 必须实际读取每个允许服务目录下的 agent.probeFileName。DaemonSet 通过同一 ConfigMap 渲染 probe_directories.txt，exec 探针逐项执行 test -r /logs/<directory>/<probeFileName>；若镜像不含 shell/test，必须随镜像提供等价的静态 probe 二进制。任一目录缺失探测文件或权限不足时不得标记 Ready。",
            "设置 seccompProfile: RuntimeDefault；资源 requests/limits 明确。",
            "不创建 ServiceAccount token 挂载：automountServiceAccountToken=false。",
        ],
    )

    add_heading(doc, "10.3 NetworkPolicy", 2)
    add_para(
        doc,
        "Agent 入站仅允许来自门户 Pod selector 的 TCP/8080。门户出站允许到 Agent Pod TCP/8080、Kubernetes API 和必要 DNS。Kubernetes API 的地址匹配依集群 CNI 与控制面部署方式配置，基础清单中应把这一差异写成 overlay，而不是放宽为任意出站。NetworkPolicy 基于 namespace 与 pod label，不是强身份；若同 namespace 存在可由非可信主体创建 Pod 的权限，必须使用独立 namespace、准入策略限制标签，或叠加 mTLS/服务网格身份。",
    )

    add_heading(doc, "11. 文件列表与路径安全", 1)
    add_heading(doc, "11.1 参数校验顺序", 2)
    add_numbers(
        doc,
        [
            "路由层只接收独立的 node、service、filename；拒绝空值和超长值。",
            "框架完成一次 URL 解码后，校验 UTF-8、控制字符、NUL、斜杠、反斜杠、点目录和残留百分号编码；任何残留 % 均拒绝，避免二次编码差异。",
            "service 只用于白名单查找，真正目录名只取配置中的 directory。",
            "filename 必须是 NFC-stable：NFC(filename)==filename，长度 1–255 字节，且不是 . 或 ..；构造上游 URL 时对 directory 和 filename 分别 PathEscape，禁止整体路径转义。",
            "上游访问依赖 Nginx disable_symlinks；文件列表仅展示普通文件类型。",
        ],
    )
    add_heading(doc, "11.2 安全组件规则", 2)
    add_table(
        doc,
        ["输入", "允许", "拒绝示例"],
        [
            ["node", "Kubernetes DNS-1123 subdomain、≤253 字节，精确匹配", "../node、node%2fetc、控制字符"],
            ["service id", "配置中精确存在", "任意目录名、大小写模糊匹配"],
            ["directory", "^[a-z0-9][a-z0-9_-]{0,62}$；启动期校验", ".、..、a/b、a\\b、A、%2f"],
            ["filename", "NFC-stable UTF-8 单一名称，1–255 字节", ".、..、a/b、a\\b、NUL、%、%252f、CR/LF"],
            ["Range", "bytes=start-end / start- / -suffix 单一区间", "其他单位、多区间、负数、溢出"],
        ],
        [1.1, 2.75, 2.65],
        compact=True,
    )
    add_para(
        doc,
        "文件列表与下载必须复用同一个 ValidateFilename 实现。列表阶段过滤掉不满足规则的普通文件并记录计数指标；下载阶段只接受前端传回的原始文件名参数，重新执行同一校验后再访问上游。不得在列表阶段展示规范化后的替代名称，也不得在下载阶段做大小写折叠、相近名称查找或 Unicode 宽松匹配。",
    )
    add_callout(
        doc,
        "双层防护",
        "门户负责逻辑白名单、解码后校验和安全转义；Nginx 负责只读挂载、固定 alias 与禁止符号链接。任何一层都不能以“另一层会拦截”为理由省略。",
        fill="FDECEC",
        accent=RED,
    )

    add_heading(doc, "12. 下载代理实现", 1)
    add_heading(doc, "12.1 HTTP Transport", 2)
    add_bullets(
        doc,
        [
            "使用独立 http.Transport：DialContext 3s、ResponseHeaderTimeout 10s、合理的 MaxIdleConnsPerHost、DisableCompression=true；下载请求不设置总 Client.Timeout，避免固定总时长截断大文件。",
            "下载复制循环必须实现 bodyIdleTimeout 进展看门狗：从上游读取或向客户端成功写出任一字节即刷新 deadline；连续 bodyIdleTimeout 无字节进展时取消上游 context、关闭 Body、释放并发槽并记录 timeout_body_idle。若配置 maxDuration>0，到达总时长也按 timeout_total 取消。",
            "每个请求绑定浏览器 Request Context；客户端断开时上游立即取消。",
            "全局带权信号量限制并发下载，获取失败快速返回 429，不让 goroutine 无限排队。",
            "连接 Agent 使用 Pod IP，禁止自动跟随重定向；任何 3xx 作为上游协议错误处理。",
        ],
    )

    add_heading(doc, "12.2 头部与正文转发", 2)
    add_code(
        doc,
        """
request allowlist  = [Range, If-Range]
forced request     = [Accept-Encoding: identity]
response allowlist = [
  Content-Type, Content-Length, Content-Range,
  Accept-Ranges, ETag, Last-Modified
]
buffer size        = 64 KiB (sync.Pool 复用)
copy               = io.CopyBuffer(ResponseWriter, upstream.Body, buffer)
""",
    )
    add_para(
        doc,
        "必须先校验上游状态和响应头，再写浏览器状态。未带 Range 时允许 200；带 Range 时只允许 206 或 416。206 的 Content-Range 必须能解析为 bytes start-end/size 且覆盖请求区间，416 必须保留 bytes */size；Content-Length 只能是非负十进制，且不得与 Content-Range 长度矛盾。Content-Disposition 由门户根据已验证文件名生成；对 ASCII fallback 去除 CR/LF、引号和路径字符。HEAD 可作为后续优化，但第一版浏览器接口仅声明 GET。",
    )

    add_heading(doc, "12.3 审计计数", 2)
    add_para(
        doc,
        "使用 countingWriter 记录实际写给客户端的字节数，而不是信任 Content-Length。事件字段 completed 仅在 io.CopyBuffer 正常结束且上游 Body 关闭成功时为 true；context canceled、broken pipe 或上游读错误均记录失败类别。",
    )

    add_heading(doc, "13. 前端页面实现", 1)
    add_heading(doc, "13.1 页面状态机", 2)
    add_code(
        doc,
        """
BOOTSTRAP
  ├─ session=401 ─► LOGIN
  └─ session=200 ─► MAIN_LOADING_NODES
MAIN_LOADING_NODES
  ├─ success ─► MAIN_SELECT_NODE
  └─ failure ─► MAIN_ERROR
MAIN_SELECT_NODE ─► MAIN_SELECT_SERVICE ─► MAIN_LOADING_FILES ─► MAIN_READY
任一 API=401 ─► 清理本地状态 ─► LOGIN
""",
    )
    add_heading(doc, "13.2 交互规则", 2)
    add_bullets(
        doc,
        [
            "节点选择器显示 Ready、NotReady、Agent 不可用；不可用节点仍可见但禁用文件加载。",
            "切换节点时清空服务和文件选择；切换服务时清空过滤值与旧文件列表。",
            "文件过滤仅在浏览器对已加载 items 做不区分大小写的包含匹配，不把任意过滤文本发送到服务器。",
            "下载使用普通链接或 window.location 触发，浏览器可接管文件保存和 Range；错误下载由服务端返回可读页面或前端预检后提示。",
            "手动刷新只刷新当前节点/服务的文件列表；加载期间禁用重复操作。",
            "所有错误显示稳定中文消息和 requestId；不展示原始 response body、内部地址或堆栈。",
        ],
    )
    add_heading(doc, "13.3 可访问性与安全", 2)
    add_bullets(
        doc,
        [
            "表单 label 与输入框显式关联，状态更新使用 aria-live，键盘可完成全部操作。",
            "文件大小使用 IEC 单位显示，时间按浏览器本地时区显示，同时保留精确 title。",
            "不把密码写入日志、localStorage 或 URL；登录成功后立即清空密码字段。",
            "默认 CSP：default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self'; connect-src 'self'; form-action 'self'; object-src 'none'; base-uri 'none'; frame-ancestors 'none'。",
        ],
    )

    add_heading(doc, "14. 审计、运行日志与指标", 1)
    add_heading(doc, "14.1 审计事件", 2)
    add_code(
        doc,
        """
{"timestamp":"2026-07-09T02:00:00Z","eventType":"login",
 "requestId":"01J...","username":"admin","clientIP":"10.0.0.8",
 "success":false,"reason":"invalid_credentials"}

{"timestamp":"2026-07-09T02:03:10Z","eventType":"download",
 "requestId":"01J...","username":"admin","clientIP":"10.0.0.8",
 "node":"worker-a","service":"order-service","filename":"app.log",
 "httpStatus":206,"bytesTransferred":1048576,"durationMs":922,
 "completed":true}
""",
    )
    add_bullets(
        doc,
        [
            "审计字段使用固定 schema 和枚举原因；不得写 password、Cookie、Authorization、文件内容、宿主机路径或 Pod IP。",
            "客户端 IP 仅在直接来源属于 trustedProxyCIDRs 时解析 Forwarded/X-Forwarded-For；否则使用 RemoteAddr。",
            "对路径攻击、无效 Range、异常 Origin 和服务白名单探测写 security 事件，保留 requestId 和分类，不回显危险原文。",
            "审计时间统一使用 UTC RFC3339Nano；平台日志系统保留期不少于 180 天并限制普通应用写入覆盖/删除权限；节点时钟漂移纳入监控，超过 2s 告警。",
        ],
    )

    add_heading(doc, "14.2 可观测性", 2)
    add_table(
        doc,
        ["类型", "建议项"],
        [
            ["健康检查", "/healthz 仅表示进程存活；/readyz 要求配置有效、informer 已同步、快照未过期"],
            ["计数器", "login_total、api_requests_total、downloads_total、security_rejections_total"],
            ["直方图", "api_duration_seconds、download_duration_seconds、upstream_header_seconds、download_body_idle_seconds"],
            ["仪表", "active_downloads、discovery_snapshot_age_seconds、ready_agents、audit_log_retention_days"],
            ["告警", "Agent 缺失、快照持续过期、登录失败突增、下载错误率、活跃下载长期满额、节点时钟漂移、审计留存不足"],
        ],
        [1.25, 5.25],
        compact=True,
    )

    add_heading(doc, "15. 错误处理矩阵", 1)
    add_table(
        doc,
        ["场景", "门户行为", "HTTP", "前端提示"],
        [
            ["会话过期", "删除 Cookie，记录 auth 事件", "401", "登录状态已失效，请重新登录"],
            ["参数/路径非法", "拒绝上游请求，写 security 事件", "400", "请求参数不合法"],
            ["服务不在白名单", "不泄露配置", "404", "服务不存在"],
            ["节点不存在", "不访问 Agent", "404", "节点不存在"],
            ["无 Ready Agent", "记录 node 与快照年龄", "503", "节点日志代理不可用"],
            ["文件轮转/删除", "映射上游 404", "404", "文件已不存在"],
            ["列表超过项目/字节上限", "立即取消上游，不完整缓冲", "413", "文件数量过多，无法展示"],
            ["Agent 权限/协议/响应异常", "隐藏上游正文和内部地址", "502", "节点日志代理响应异常"],
            ["上游连接/响应头/正文空转超时", "取消上游，释放并发槽", "504", "节点响应超时，请稍后重试"],
            ["客户端中断下载", "取消上游，completed=false", "连接中断", "由浏览器处理，可再次下载"],
            ["Kubernetes 缓存过期", "拒绝新解析", "503", "集群状态暂时不可用"],
        ],
        [1.55, 2.5, 0.75, 1.7],
        compact=True,
    )

    add_heading(doc, "16. Kubernetes 部署资源", 1)
    add_heading(doc, "16.1 资源清单", 2)
    add_table(
        doc,
        ["资源", "关键要求"],
        [
            ["Portal Deployment", "replicas=1；只读根文件系统；非 root；配置与 Secret 分挂载；探针和资源限制"],
            ["Portal Service", "ClusterIP，仅暴露门户 HTTP 端口"],
            ["Ingress", "强制 HTTPS；关闭响应缓冲和临时文件落盘；读写超时大于 download.bodyIdleTimeout 且覆盖 maxDuration；保留 Range 响应；配置按控制器 overlay 固化"],
            ["ServiceAccount/RBAC", "Node get/list/watch；指定 namespace 的 Agent Pod get/list/watch"],
            ["Agent DaemonSet", "每个目标 Node 一个 Pod；hostPath ro；固定 UID/GID 或 supplementalGroups 可读；exec/二进制权限探测；容忍度与 nodeSelector 按环境 overlay"],
            ["Agent ConfigMap", "Nginx 配置、probe_directories.txt 与 probeFileName；只开启 /logs/ 的 GET/HEAD 与 JSON autoindex"],
            ["NetworkPolicy", "Agent 仅收 Portal；Portal 仅出站到 Agent、K8s API、DNS；命名空间/标签权限受控"],
            ["Secret", "username、password hash、session key；禁止提交真实值"],
        ],
        [1.65, 4.85],
        compact=True,
    )

    add_heading(doc, "16.2 RBAC 基线", 2)
    add_code(
        doc,
        """
rules:
  - apiGroups: [""]
    resources: ["nodes"]
    verbs: ["get", "list", "watch"]
  - apiGroups: [""]
    resources: ["pods"]
    verbs: ["get", "list", "watch"]
    # Role 绑定到 Agent 所在 namespace
""",
    )
    add_para(
        doc,
        "Node 是集群级资源，使用 ClusterRole；Pod 可使用 namespace Role。若平台要求统一 ClusterRole，仍应通过 namespace scope 和固定 label selector 在客户端进一步收窄。",
    )

    add_heading(doc, "16.3 发布与回滚", 2)
    add_numbers(
        doc,
        [
            "先发布 Agent ConfigMap 与 DaemonSet，确认 probe_directories.txt 与门户 services[].directory 同源生成；目标 Node 全部有 Agent Pod，挂载只读，非 root 用户可遍历每个允许目录并读取 agent.probeFileName，权限不足时 readiness=false；NetworkPolicy 生效。",
            "再发布 Portal 的 ServiceAccount/RBAC、ConfigMap、Secret、Deployment 与 Service。",
            "最后开放 Ingress，确认响应缓冲和临时文件落盘已关闭，读写超时与 bodyIdleTimeout/maxDuration 一致，并从集群外执行登录、列表、1 GiB+ 完整下载、Range、上游空转、路径攻击和审计冒烟测试。",
            "回滚 Portal 镜像不会影响 Agent 和宿主机日志；会话会失效，用户重新登录。",
            "回滚 Agent 前确认门户与 Nginx 配置兼容；滚动期间节点可短暂显示 Agent 不可用。",
        ],
    )

    add_heading(doc, "17. 测试方案", 1)
    add_heading(doc, "17.1 单元测试", 2)
    add_table(
        doc,
        ["模块", "必须覆盖"],
        [
            ["config", "合法配置、重复 ID、directory 正则、probeFileName、边界时长、密钥缺失"],
            ["security", "一次/二次编码、斜杠与反斜杠、点目录、NUL、控制字符、NFC-stable、残留 %、超长输入"],
            ["auth", "正确/错误密码、虚拟校验、Cookie 篡改、过期、退出、限流窗口"],
            ["discovery", "Node Ready、Pod Ready、终止 Pod、多 Pod、IPv4/IPv6、DNS subdomain、稳定集群无事件、BOOKMARK/事件推进 RV、watch 建立但未推进 RV、relist、快照过期"],
            ["agent", "autoindex JSON 流式解码、响应体字节上限、非文件项、坏 mtime、负 size、第 10,001 项提前终止"],
            ["download", "200、206、416、Range 请求返回 200、Content-Range/Content-Length 校验、bodyIdleTimeout、单 Range、非法/多 Range、Accept-Encoding identity、禁用自动解压、头部白名单、取消、计数字节"],
            ["api", "状态映射、requestId 接收/拒绝/生成、错误 envelope、无内部信息泄漏"],
            ["audit", "字段完整、原因枚举、UTC 时间、敏感字段扫描、保留策略配置检查"],
        ],
        [1.25, 5.25],
        compact=True,
    )

    add_heading(doc, "17.2 集成与端到端测试", 2)
    add_bullets(
        doc,
        [
            "使用 fake clientset 或 envtest 驱动 Node/Pod 变化，验证 informer 快照和 API 输出。",
            "启动真实测试 Nginx，挂载临时只读目录，验证 JSON 列表、完整下载、Range、ETag 与轮转后 404。",
            "模拟慢响应头、上游正文连续无字节、上游读中断、客户端取消、Range 请求返回 200 和 3xx，验证超时、资源释放与审计 completed。",
            "在至少两个 Node 的测试集群部署，验证每个 Agent 只能读取所在节点目录，且缺失 agent.probeFileName 或权限不足时 readiness=false。",
            "从非门户 Pod 访问 Agent 应失败；从浏览器/集群外访问 Agent 无路由；在可创建 Pod 的非可信命名空间验证不能通过标签冒充 Portal。",
            "创建大于 1 GiB 的稀疏测试文件，监控 Portal RSS，验证内存不随文件大小线性增长。",
        ],
    )

    add_heading(doc, "17.3 关键安全用例", 2)
    add_code(
        doc,
        """
../secret
..%2fsecret
%252e%252e%252fsecret
file%2flog
file%5clog
file%00.log
file%0d%0aX-Test:1
Range: bytes=0-10,20-30
Range: items=0-10
X-Request-ID: bad id with spaces
X-Request-ID: <script>alert(1)</script>
""",
    )
    add_para(
        doc,
        "每个样例都必须在门户层拒绝，不能向 Nginx 发起请求；测试同时断言响应无内部路径、Pod IP 和原始输入回显，并产生 security 审计事件。",
    )

    add_heading(doc, "18. 开发任务与交付顺序", 1)
    add_para(
        doc,
        "按可独立验证的纵向切片实施。每个任务遵循“先失败测试、最小实现、测试通过、提交”的节奏；部署清单与代码同步更新。",
    )
    add_table(
        doc,
        ["阶段", "交付物", "完成判据"],
        [
            ["1. 工程骨架", "Go module、main、配置、健康检查、CI", "非法配置失败；healthz/readyz 测试通过"],
            ["2. 认证会话", "密码校验、内存会话、Cookie、限流、登录 API", "登录/过期/篡改/限流测试通过"],
            ["3. K8s 发现", "informer、快照、Node/Pod resolver、节点 API", "多 Pod、BOOKMARK/RV、缓存过期、IPv6 测试通过"],
            ["4. Agent 列表", "Nginx 配置、DaemonSet、autoindex client、文件 API", "真实 Nginx 集成测试通过"],
            ["5. 下载代理", "Range、头白名单、禁用自动解压、流式复制、正文空转看门狗、并发限制、审计", "1 GiB+ 内存、中断与空转测试通过"],
            ["6. 前端", "登录、选择器、文件表、过滤、刷新、错误态", "键盘操作、401 跳转、错误 requestId 验证通过"],
            ["7. 安全部署", "RBAC、NetworkPolicy、Ingress、安全上下文、宿主机读取权限", "跨 Pod访问、标签冒充、权限探测、集群外 1 GiB+ Range 验证通过"],
            ["8. 验收", "E2E、攻击用例、审计、运维说明", "第 19 节矩阵全部通过"],
        ],
        [1.15, 3.1, 2.25],
        compact=True,
    )

    add_heading(doc, "18.1 建议提交边界", 2)
    add_bullets(
        doc,
        [
            "每个阶段至少一个可回滚提交，不把前端、RBAC 和下载核心混在同一提交。",
            "测试与实现同提交；安全回归样例必须在修复提交中落库。",
            "真实 Secret、集群地址和宿主机日志不得进入 Git；示例值统一使用占位标识。",
            "合并前执行 go test ./...、静态分析、容器镜像扫描、Kubernetes 清单校验和集成冒烟。",
        ],
    )

    add_heading(doc, "19. 验收追踪矩阵", 1)
    add_table(
        doc,
        ["设计验收项", "实现位置", "验证方式"],
        [
            ["统一 Web 登录入口", "auth + api + webui + Ingress", "E2E 登录/退出/过期"],
            ["按节点、服务、文件下载", "discovery + agent + download + UI", "两节点端到端下载"],
            ["统一目录规则", "ConfigMap + Agent hostPath", "部署清单与挂载检查"],
            ["1 GiB+ 流式下载、内存稳定", "download.proxy io.CopyBuffer + Ingress overlay", "集群外 HTTPS 下载、RSS 曲线、Ingress 无临时文件"],
            ["HTTP Range 续传", "download.range + Nginx", "206/Content-Range/断点续传"],
            ["下载正文空转不会占死资源", "download.bodyIdleTimeout + context cancel", "上游/下游无字节进展超时、并发槽释放"],
            ["发现缓存新鲜度严谨", "informer RV/BOOKMARK/relist", "watch 建立不推进 RV 不刷新，relist 恢复"],
            ["白名单外不可访问", "config lookup + security component", "404 与无上游请求断言"],
            ["路径编码/符号链接不可逃逸", "security + disable_symlinks + ro mount", "攻击语料与集群验证"],
            ["Agent 权限探针可执行", "Agent ConfigMap + readiness probe", "缺失探针文件或权限不足时 NotReady"],
            ["单节点故障不影响其他节点", "resolver 按节点隔离", "停止一个 Agent 后验证其他节点"],
            ["登录与下载审计完整无敏感信息", "audit schema + countingWriter", "schema 断言、UTC 时间、留存与敏感词扫描"],
            ["无业务数据库", "内存会话 + stdout 审计", "部署资源与依赖检查"],
        ],
        [2.25, 2.35, 1.9],
        compact=True,
    )

    add_heading(doc, "20. 风险与演进边界", 1)
    add_table(
        doc,
        ["风险", "第一版控制", "触发后续演进的信号"],
        [
            ["中央门户成为带宽瓶颈", "并发限制、流式代理、指标", "网卡持续饱和或排队显著，评估短期签名直链"],
            ["单副本导致短时不可用", "快速重启、无业务数据、会话可重登", "可用性目标提高，改造共享会话或无状态票据"],
            ["Nginx 元数据能力有限", "JSON autoindex + disable_symlinks", "需要更细鉴权/哈希/递归时评估 Go Agent"],
            ["共享管理员缺少责任区分", "完整登录与下载审计", "用户规模扩大时接入 OIDC/SSO 与 RBAC"],
            ["目录中文件数量过大", "流式 JSON、响应体字节上限、10,000 项提前终止", "需要分页/索引时引入 Agent API，不直接放宽限制"],
        ],
        [1.65, 2.5, 2.35],
        compact=True,
    )

    add_callout(
        doc,
        "禁止提前扩展",
        "多用户、文件预览、Tail、多集群、打包下载和签名直链不进入第一版。任何新增能力必须先更新威胁模型、接口契约与验收矩阵。",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "附录 A：发布前检查清单", 1)
    add_bullets(
        doc,
        [
            "配置与 Secret 校验通过，仓库和镜像中无明文凭据。",
            "Portal 与 Agent 均非 root、只读根文件系统、禁止提权、drop all capabilities。",
            "hostPath readOnly，非 root Agent 的 UID/GID 或 supplementalGroups 可遍历目录并读取日志；每个允许目录存在 agent.probeFileName 且 readiness 探针验证通过；Agent 不挂载 ServiceAccount token。",
            "RBAC 仅含 Node 与指定 namespace Pod 的 get/list/watch。",
            "Agent 无 Ingress/NodePort；NetworkPolicy 实测只允许 Portal，且命名空间/标签冒充风险已通过准入或隔离控制。",
            "HTTPS、Cookie Secure/HttpOnly/SameSite、CSP 与 requestId 已验证。",
            "路径攻击语料、符号链接、非法 Range 和响应头注入测试通过。",
            "经集群外 HTTPS Ingress 的 1 GiB+ 大文件、断点续传、无缓冲落盘、正文空转、客户端取消、节点故障、K8s API 故障测试通过。",
            "登录与下载审计字段完整，敏感信息扫描为零；日志留存、不可覆盖/删除权限和节点时钟漂移告警已验证。",
            "告警、回滚步骤和运行手册已交付给运维人员。",
        ],
    )

    add_heading(doc, "附录 B：术语", 1)
    add_table(
        doc,
        ["术语", "说明"],
        [
            ["Portal", "面向用户的 Go 中央门户"],
            ["Agent", "每个 Node 上只读访问日志目录的 Nginx Pod"],
            ["逻辑服务", "ConfigMap 中以 id 标识并映射到 directory 的可访问服务"],
            ["发现快照", "由 informer 生成的 Node 与 Ready Agent Pod 不可变映射"],
            ["普通文件", "可由 Nginx 静态读取、非目录且不经符号链接逃逸的文件"],
            ["完成下载", "上游读取与向客户端复制均正常结束的请求"],
        ],
        [1.35, 5.15],
    )

    doc.core_properties.title = "Kubernetes 节点日志下载门户开发文档"
    doc.core_properties.subject = "Go 中央门户与 Nginx DaemonSet 的开发实施基线"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Kubernetes, Go, Nginx, 日志下载, 开发文档"
    doc.core_properties.comments = "根据 2026-07-09 设计文档生成；v1.1-r2 修复设计/开发文档评审问题"
    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
